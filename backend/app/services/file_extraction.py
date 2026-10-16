import csv
import io
import re
from datetime import date as date_type
from html.parser import HTMLParser
from pathlib import Path

MAIL_EXTENSIONS = (".eml", ".msg")


def extract_text(file_path: Path, original_name: str) -> str:
    suffix = Path(original_name).suffix.lower()

    try:
        if suffix == ".pdf":
            return _extract_pdf(file_path)
        if suffix == ".docx":
            return _extract_docx(file_path)
        if suffix in (".xlsx", ".xlsm"):
            return _extract_xlsx(file_path)
        if suffix == ".csv":
            return _extract_csv(file_path)
        if suffix in (".txt", ".md"):
            return file_path.read_text(encoding="utf-8", errors="ignore")
        if suffix in MAIL_EXTENSIONS:
            return eml_to_text(parse_mail(file_path))
        if suffix == ".pptx":
            return _extract_pptx(file_path)
    except Exception as exc:
        return f"[שגיאה בחילוץ טקסט אוטומטי מהקובץ: {exc}]"

    return "[לא ניתן לחלץ טקסט אוטומטית מסוג קובץ זה. ניתן להוסיף תוכן ידנית.]"


def _extract_pdf(file_path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    parts = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"--- עמוד {i + 1} ---\n{text.strip()}")
    return "\n\n".join(parts)


def _extract_docx(file_path: Path) -> str:
    import docx

    doc = docx.Document(str(file_path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(file_path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"--- גיליון: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_csv(file_path: Path) -> str:
    parts = []
    with open(file_path, "r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            parts.append("\t".join(row))
    return "\n".join(parts)


class _HtmlText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("br", "p", "div", "tr", "li"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _html_to_text(html: str) -> str:
    parser = _HtmlText()
    parser.feed(html)
    text = "".join(parser.parts)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg")


def parse_eml(file_path: Path, with_attachments: bool = False) -> dict:
    """Headers and body of an .eml file: {subject, date, from, to, cc, body, attachments, attachment_files,
    attachment_skipped}; people are (name, email) pairs. attachment_files ({name, data}) is filled only on request."""
    from email import policy
    from email.parser import BytesParser
    from email.utils import getaddresses, parsedate_to_datetime

    with open(file_path, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)

    def people(header: str) -> list[tuple[str, str]]:
        return [(str(name or "").strip(), str(addr or "").strip()) for name, addr in getaddresses([str(v) for v in msg.get_all(header, [])]) if addr]

    sent = None
    try:
        sent = parsedate_to_datetime(str(msg["date"])).date() if msg["date"] else None
    except (TypeError, ValueError):
        sent = None

    body = ""
    part = msg.get_body(preferencelist=("plain", "html"))
    if part is not None:
        content = part.get_content()
        body = _html_to_text(content) if part.get_content_type() == "text/html" else str(content).strip()

    files, skipped = [], []
    if with_attachments:
        for part in msg.iter_attachments():
            ctype = part.get_content_type()
            if part.get("Content-ID") and ctype.startswith("image/"):
                continue  # a picture embedded in the text (signature, logo), not a real attachment
            name = part.get_filename() or ""
            if ctype == "message/rfc822":  # a mail attached as a message: keep it as an .eml file
                inner = part.get_payload()
                data = inner[0].as_bytes() if inner else None
                name = name or "message.eml"
            else:
                data = part.get_payload(decode=True)
            if not data:
                skipped.append(name or "(ללא שם)")
                continue
            files.append({"name": Path(name or "attachment").name, "data": data})

    return {
        "subject": str(msg["subject"] or "").strip(),
        "date": sent if isinstance(sent, date_type) else None,
        "from": people("from"),
        "to": people("to"),
        "cc": people("cc"),
        "body": body,
        "attachments": [a.get_filename() or "(ללא שם)" for a in msg.iter_attachments()],
        "attachment_files": files,
        "attachment_skipped": skipped,
    }


def _format_people(pairs: list[tuple[str, str]]) -> str:
    return ", ".join(f"{name} <{addr}>" if name else addr for name, addr in pairs)


def eml_to_text(mail: dict, mention_attachments: bool = True) -> str:
    lines = []
    if mail["from"]:
        lines.append(f"מאת: {_format_people(mail['from'])}")
    if mail["to"]:
        lines.append(f"אל: {_format_people(mail['to'])}")
    if mail["cc"]:
        lines.append(f"עותק: {_format_people(mail['cc'])}")
    if mail["date"]:
        lines.append(f"תאריך: {mail['date']}")
    lines.append(f"נושא: {mail['subject']}")
    if mention_attachments and mail["attachments"]:
        lines.append("קבצים מצורפים (לא נקראו): " + ", ".join(mail["attachments"]))
    return "\n".join(lines) + "\n\n" + mail["body"]


def parse_mail(file_path: Path, with_attachments: bool = False) -> dict:
    """An .eml or Outlook .msg file as {subject, date, from, to, cc, body, attachments, attachment_files, attachment_skipped}."""
    if file_path.suffix.lower() == ".msg":
        return parse_msg(file_path, with_attachments)
    return parse_eml(file_path, with_attachments)


def _split_address(value: str) -> tuple[str, str]:
    from email.utils import parseaddr

    name, addr = parseaddr(value or "")
    return (name or "").strip(), (addr or "").strip()


def parse_msg(file_path: Path, with_attachments: bool = False) -> dict:
    """An Outlook .msg file, in the same shape as parse_eml. Recipients without an SMTP address are left out
    (Exchange-internal addresses are not usable as an e-mail address)."""
    import extract_msg

    msg = extract_msg.openMsg(str(file_path))
    try:
        sender = []
        if msg.sender:
            name, addr = _split_address(str(msg.sender))
            if addr:
                sender.append((name, addr))
        to, cc = [], []
        for r in msg.recipients or []:
            addr = (getattr(r, "smtpAddress", None) or getattr(r, "email", None) or "").strip()
            if not addr or "@" not in addr:
                continue
            name = (getattr(r, "name", None) or "").strip()
            kind = str(getattr(getattr(r, "type", None), "name", getattr(r, "type", ""))).lower()
            (cc if "cc" in kind and "bcc" not in kind else to).append((name, addr))
        sent = msg.date
        body = (msg.body or "").strip()
        if not body and msg.htmlBody:
            html = msg.htmlBody.decode("utf-8", errors="ignore") if isinstance(msg.htmlBody, (bytes, bytearray)) else str(msg.htmlBody)
            body = _html_to_text(html)
        attachments, files, skipped = [], [], []
        for a in msg.attachments or []:
            name = getattr(a, "name", None) or getattr(a, "longFilename", None) or getattr(a, "shortFilename", None) or "(ללא שם)"
            attachments.append(name)
            if not with_attachments or getattr(a, "hidden", False):
                continue
            if getattr(a, "cid", None) and Path(name).suffix.lower() in IMAGE_EXTENSIONS:
                continue  # a picture embedded in the text (signature, logo), not a real attachment
            data = getattr(a, "data", None)
            if isinstance(data, (bytes, bytearray)) and data:
                files.append({"name": Path(name).name, "data": bytes(data)})
            else:  # an attached message or something the library cannot hand over as bytes
                skipped.append(name)
        return {
            "subject": str(msg.subject or "").strip(),
            "date": sent.date() if hasattr(sent, "date") else None,
            "from": sender,
            "to": to,
            "cc": cc,
            "body": body,
            "attachments": attachments,
            "attachment_files": files,
            "attachment_skipped": skipped,
        }
    finally:
        msg.close()


def _extract_pptx(file_path: Path) -> str:
    from pptx import Presentation

    parts = []
    for i, slide in enumerate(Presentation(str(file_path)).slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[הערות מרצה] {notes}")
        if lines:
            parts.append(f"--- שקף {i} ---\n" + "\n".join(lines))
    return "\n\n".join(parts)
